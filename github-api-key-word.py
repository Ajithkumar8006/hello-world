from openai import OpenAI
from datetime import datetime
from docx import Document
import os
import traceback

GITHUB_TOKEN = os.getenv("GITHUB_TOKEN")

if not GITHUB_TOKEN:
    raise ValueError("GITHUB_TOKEN environment variable is not set")

PROMPT_FILE = "rca-prompt.txt"
INCIDENT_FILE = "azure_incident_report_20260522_133208.txt"

MODEL = "openai/gpt-4.1-mini"

OUTPUT_FILE = f"RCA_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"

client = OpenAI(
    api_key=GITHUB_TOKEN,
    base_url="https://models.github.ai/inference"
)

def read_file(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()

try:
    rca_prompt = read_file(PROMPT_FILE)
    incident_report = read_file(INCIDENT_FILE)

    full_prompt = f"""
{rca_prompt}

====================================================
INCIDENT REPORT
====================================================

{incident_report}
"""

    response = client.chat.completions.create(
        model=MODEL,
        temperature=0.2,
        messages=[
            {
                "role": "system",
                "content": "You are a Senior Azure Cloud Reliability Engineer and Incident Manager."
            },
            {
                "role": "user",
                "content": full_prompt
            }
        ]
    )

    rca_report = response.choices[0].message.content

    if not rca_report:
        raise Exception("Empty response received from model")

    doc = Document()

    doc.add_heading("Root Cause Analysis (RCA)", level=1)

    for line in rca_report.splitlines():
        doc.add_paragraph(line)

    doc.save(OUTPUT_FILE)

    print(f"RCA report saved: {OUTPUT_FILE}")

except Exception:
    traceback.print_exc()

