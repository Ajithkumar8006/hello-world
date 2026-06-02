import os
import json
import requests
from docx import Document

# OpenRouter API Key
api_key = os.getenv("OPENROUTER_API_KEY")

if not api_key:
    raise Exception(
        "OPENROUTER_API_KEY environment variable not set"
    )

# Read RCA Prompt Instructions
with open("rca-prompt.txt", "r", encoding="utf-8") as f:
    rca_prompt = f.read().strip()

# Read Incident Report
with open(
    "azure_incident_report_20260522_133208.txt",
    "r",
    encoding="utf-8"
) as f:
    incident_report = f.read()

# Build AI Prompt
prompt = f"""
{rca_prompt}

Review the following Azure incident report.

Generate a professional Root Cause Analysis (RCA) document.

Incident Report:

{incident_report}
"""

print("Sending report to OpenRouter...")

response = requests.post(
    "https://openrouter.ai/api/v1/chat/completions",
    headers={
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json"
    },
    json={
        "model": "openai/gpt-oss-120b:free",
        "messages": [
            {
                "role": "user",
                "content": prompt
            }
        ]
    },
    timeout=300
)

print("Status Code:", response.status_code)

result = response.json()

if "choices" not in result:
    print(json.dumps(result, indent=2))
    raise Exception("OpenRouter request failed")

answer = result["choices"][0]["message"]["content"]

# Create DOCX
doc = Document()

doc.add_heading(
    "Azure Incident RCA Review",
    level=1
)

for line in answer.splitlines():
    if line.strip():
        doc.add_paragraph(line)

output_file = "incident_rca_review.docx"

doc.save(output_file)

print(f"Created {output_file}")
